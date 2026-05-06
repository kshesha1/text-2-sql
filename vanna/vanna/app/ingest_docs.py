"""Ingest PDFs / PowerPoint decks into vanna's pgvector 'documentation' collection.

Each file is extracted to plain text, split into fixed-size character chunks
with a small overlap, and pushed via ``vn.add_documentation``. The chunks land
in the ``documentation`` PGVector collection alongside any business glossary
text, and vanna will retrieve them as RAG context for SQL generation.
"""

from __future__ import annotations

from pathlib import Path
from typing import Iterable, List

# pyrefly: ignore [missing-import]
from pypdf import PdfReader
# pyrefly: ignore [missing-import]
from pptx import Presentation


def extract_pdf(path: Path) -> str:
    reader = PdfReader(str(path))
    return "\n".join((page.extract_text() or "") for page in reader.pages)


def extract_pptx(path: Path) -> str:
    out: List[str] = []
    prs = Presentation(str(path))
    for slide_idx, slide in enumerate(prs.slides, start=1):
        out.append(f"--- Slide {slide_idx} ---")
        for shape in slide.shapes:
            text = getattr(shape, "text", None)
            if text:
                out.append(text)
    return "\n".join(out)


def chunk_text(text: str, size: int = 1500, overlap: int = 200) -> List[str]:
    if size <= 0:
        raise ValueError("size must be positive")
    if overlap >= size:
        raise ValueError("overlap must be smaller than size")
    chunks: List[str] = []
    i = 0
    n = len(text)
    while i < n:
        chunks.append(text[i : i + size])
        i += size - overlap
    return chunks


def _extract(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return extract_pdf(path)
    if suffix in (".pptx", ".ppt"):
        return extract_pptx(path)
    if suffix in (".txt", ".md"):
        return path.read_text(encoding="utf-8", errors="ignore")
    raise ValueError(f"Unsupported file type: {path.suffix} ({path})")


def ingest(vn, paths: Iterable[str | Path]) -> int:
    """Extract → chunk → add_documentation. Returns total chunks ingested."""
    total = 0
    for raw in paths:
        path = Path(raw)
        if not path.exists():
            print(f"[skip] {path} does not exist")
            continue
        text = _extract(path)
        if not text.strip():
            print(f"[skip] {path} produced no extractable text")
            continue
        chunks = chunk_text(text)
        for chunk in chunks:
            vn.add_documentation(f"[source: {path.name}]\n{chunk}")
        print(f"[ok] {path.name}: {len(chunks)} chunks")
        total += len(chunks)
    return total
